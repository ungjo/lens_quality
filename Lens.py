from pathlib import Path    #입력 Excel, DB, 결과 파일의 경로 지정.
from contextlib import closing
from dataclasses import dataclass
#데이터베이스 사용을 위한 소프트웨어 import. SQLite 데이터베이스에 연결하고 SQL 실행
#sqlite : MySQL, Oracle과 다르게 서버가 없어도 라이브러리 형태로 실행 가능. 모든 데이터가 하나의 파일에 다 저장됨. 설정 없이 즉시 사용 가능.
import sqlite3
from tempfile import TemporaryDirectory
from typing import Callable

import pandas as pd #Pandas 라이브러리를 pd라는 이름으로 import -> 표 형태의 데이터를 다루기 위해 사용
from openpyxl.styles import PatternFill #Excel 셀의 배경색 지정을 위해 openpyxl.styles의 PatternFill을 import


# 파일 위치: 이 Python 파일이 있는 폴더를 기준으로 설정
BASE = Path(__file__).resolve().parent

INPUT = BASE / "data" / "lens_qualityCheck.xlsx"

RESULTS = BASE.parents[1] / "resources" / "lens_quality" / "outputs" / "results"
DB_FILE = RESULTS / "lens_quality.db"
OUTPUT = RESULTS / "lens_inspection_result.xlsx"

# 기준표에서 항목 전체가 빠져도 검사 누락을 발견할 수 있도록 목록을 따로 둡니다.
REQUIRED_ITEMS = (
    "스크래치등급", "기포이물질개수", "코팅불량개수", "굴절률",
    "광투과율", "왜곡률", "렌즈직경", "중심두께", "가장자리두께",
    "곡률반경", "UV차단율", "색상균일도", "틴트농도",
)


@dataclass
class InspectionReport:
    """터미널과 GUI가 함께 사용하는 검사 결과. 판정과 파일 저장 상태를 구분합니다."""
    """결과표, 항목별 상세, 합격/불합격 개수, 저장 경로 등을 GUI와 터미널 모두에서 사용"""

    measurements: pd.DataFrame
    criteria: pd.DataFrame
    result: pd.DataFrame
    detail: pd.DataFrame
    counts: dict[str, int]
    input_path: Path
    db_path: Path
    excel_path: Path
    excel_saved: bool = False
    excel_error: str | None = None


def check_value(item, value, lower, upper):
    """항목 하나의 측정값을 하한·상한과 비교"""
    if pd.isna(value):
        return "판정불가", f"{item}: 측정값 누락"

    if pd.isna(lower) or pd.isna(upper):
        return "판정불가", f"{item}: 기준 누락"

    if lower > upper:
        return "판정불가", f"{item}: 하한이 상한보다 큼"

    if value < lower:
        return "불합격", f"{item}: {value:g} < 하한 {lower:g}"

    if value > upper:
        return "불합격", f"{item}: {value:g} > 상한 {upper:g}"

    return "합격", f"{item}: 허용범위 {lower:g}~{upper:g} 이내"


def format_inspection_results(writer, result, detail):
    """판정 색상과 불합격 원인 측정값을 표시합니다. 재실행 시 갱신"""
    fills = {
        "합격": PatternFill(fill_type="solid", fgColor="D9EAD3"),
        "불합격": PatternFill(fill_type="solid", fgColor="F4CCCC"),
    }
    failed_fill = PatternFill(fill_type="solid", fgColor="D9D9D9")
    summary_sheet = writer.sheets["검사결과"]
    detail_sheet = writer.sheets["항목별상세"]

    # 열 위치가 바뀌어도 열 이름과 제품번호로 올바른 셀을 찾습니다.
    summary_columns = {name: i for i, name in enumerate(result.columns, 1)}
    detail_columns = {name: i for i, name in enumerate(detail.columns, 1)}
    product_rows = {}
    for row_number, record in enumerate(result.to_dict("records"), 2):
        product_rows[record["제품번호"]] = row_number
        fill = fills.get(record["판정"])
        if fill is not None:
            summary_sheet.cell(row_number, summary_columns["판정"]).fill = fill

    for row_number, record in enumerate(detail.to_dict("records"), 2):
        status = record["항목판정"]
        fill = fills.get(status)
        if fill is not None:
            detail_sheet.cell(row_number, detail_columns["항목판정"]).fill = fill
        if status == "불합격":
            summary_sheet.cell(
                product_rows[record["제품번호"]],
                summary_columns[record["검사항목"]],
            ).fill = failed_fill
            detail_sheet.cell(row_number, detail_columns["측정값"]).fill = failed_fill

def load_excel_data(input_path):
    """Excel을 읽고 입력을 확인합니다. DB나 결과 파일은 생성하지 않습니다."""
    # 1. Excel 읽기
    # A:P에는 제품 정보와 측정값이 존재.
    # Excel의 기존 판정·판정사유 열은 읽지 않음.
    # (pandas)pd.read_excel() : Excel 파일을 읽어 DataFrame이라는 행과 열이 있는 표 형태의 데이터 구조로 변환.
    measurements = pd.read_excel(
        input_path,
        sheet_name="측정데이터",
        usecols="A:P",
    )
    # 현재 기준표의 첫 15행: 공통 12개 + 소재별 굴절률 3개
    criteria = pd.read_excel(
        input_path,
        sheet_name="기준표",
        usecols="A:F",
        nrows=15,
    )

    required_columns = {
        "측정데이터": (measurements, ["제품번호", "소재종류", "제품명", *REQUIRED_ITEMS]),
        "기준표": (criteria, ["검사항목", "적용소재", "단위", "하한", "목표치", "상한"]),
    }
    for sheet, (frame, columns) in required_columns.items():
        missing = [name for name in columns if name not in frame.columns]
        if missing:
            raise ValueError(f"{sheet} 시트에 필요한 컬럼이 없습니다: {', '.join(missing)}")

    measurements = measurements.dropna(subset=["제품번호"]) #dropna : 결측치(NaN, Null, 빈 값)가 포함된 행이나 열을 제거할 때 사용하는 함수.
    if measurements.empty:
        raise ValueError("검사할 제품이 없습니다. 측정데이터 시트의 제품번호를 확인하세요.")

    if measurements["제품번호"].duplicated().any():
        raise ValueError("제품번호가 중복되어 있습니다.")

    if criteria.duplicated(["검사항목", "적용소재"]).any():
        raise ValueError("같은 항목·소재의 기준이 중복되어 있습니다.")

    if criteria["검사항목"].isna().any():
        raise ValueError("기준표에 검사항목 이름이 비어 있는 행이 있습니다.")
    missing_items = [item for item in REQUIRED_ITEMS if item not in criteria["검사항목"].values]
    if missing_items:
        raise ValueError(f"기준표에 필수 검사항목이 없습니다: {', '.join(missing_items)}")

    # 측정값과 기준을 숫자로 변환합니다.
    # 잘못된 문자열은 자동으로 0으로 바꾸지 않고 오류를 냅니다.
    items = criteria["검사항목"].drop_duplicates().tolist()

    for item in items:
        if item not in measurements.columns:
            raise ValueError(f"측정데이터 시트에 검사항목 컬럼이 없습니다: {item}")
        try:
            measurements[item] = pd.to_numeric(measurements[item], errors="raise")
        except (TypeError, ValueError) as exc:
            raise ValueError(f"측정데이터의 '{item}'에 숫자로 변환할 수 없는 값이 있습니다.") from exc

    for column in ["하한", "목표치", "상한"]:
        try:
            criteria[column] = pd.to_numeric(criteria[column], errors="raise")
        except (TypeError, ValueError) as exc:
            raise ValueError(f"기준표의 '{column}'에 숫자로 변환할 수 없는 값이 있습니다.") from exc

    return measurements, criteria, items


def save_result_excel(report: InspectionReport, output_path=None) -> bool:
    """검사 결과를 Excel로 저장하거나 다시 저장합니다. DB 검사는 반복하지 않습니다.

    저장 성공 여부를 반환하며 report.excel_saved/excel_error에도 기록합니다.
    파일 잠금·권한 등 OSError는 실패 상태로 전달하고 다른 예외는 호출자에게 전달합니다.
    """
    output_path = Path(output_path if output_path is not None else report.excel_path).resolve()
    if output_path in (report.input_path, report.db_path):
        raise ValueError("결과 Excel은 원본 Excel이나 DB와 다른 경로에 저장해야 합니다.")

    report.excel_path = output_path
    report.excel_saved = False
    report.excel_error = None
    try:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        # 완성된 파일만 교체하여 저장 도중 오류가 나도 이전 결과를 보존합니다.
        with TemporaryDirectory(prefix=".lens_excel_", dir=output_path.parent) as temp_dir:
            temporary_file = Path(temp_dir) / output_path.name
            with pd.ExcelWriter(temporary_file, engine="openpyxl") as writer:
                report.result.to_excel(writer, sheet_name="검사결과", index=False)
                report.detail.to_excel(writer, sheet_name="항목별상세", index=False)
                report.criteria.to_excel(writer, sheet_name="적용기준", index=False)
                format_inspection_results(writer, report.result, report.detail)
            temporary_file.replace(output_path)
    except OSError as exc:
        report.excel_error = str(exc)
        return False

    report.excel_saved = True
    return True


def run_inspection(
    input_path=INPUT,
    output_dir=RESULTS,
    progress_callback: Callable[[str], None] | None = None,
) -> InspectionReport:
    """선택한 경로로 검사하고 결과를 반환합니다. 화면이나 터미널에는 직접 출력하지 않습니다.

    입력·DB 오류는 예외로 전달합니다. DB 저장 후 Excel의 I/O 오류가 나면
    excel_saved=False인 결과를 반환하여 GUI에서 안내하고 저장을 재시도할 수 있습니다.
    progress_callback은 이 함수를 실행한 스레드에서 호출됩니다. 향후 GUI에서는
    화면을 직접 바꾸는 대신 Qt 시그널을 전달하는 함수를 연결하세요.
    """
    input_path = Path(input_path).resolve() #resolve() : pathlib모듈의 Path 객체에서 제공하는 메서드. 상대경로를 절대경로로 변환
    output_dir = Path(output_dir).resolve()
    db_path = (output_dir / DB_FILE.name).resolve()
    excel_path = (output_dir / OUTPUT.name).resolve()
    if input_path in (db_path, excel_path):
        raise ValueError("원본 파일과 결과 파일의 저장 경로가 같습니다. 다른 결과 폴더를 선택하세요.")
    if not input_path.is_file():
        raise FileNotFoundError(f"원본 Excel 파일을 찾을 수 없습니다: {input_path}")

    def notify(message):
        if progress_callback is not None:
            progress_callback(message)

    notify("Excel 읽기 및 입력 확인")
    measurements, criteria, items = load_excel_data(input_path)
    output_dir.mkdir(parents=True, exist_ok=True)

    notify("입력 데이터 DB 저장")
    # 연결을 만든 스레드 안에서 사용하고, 성공·실패 모두에서 닫습니다.
    with closing(sqlite3.connect(db_path)) as conn, conn:  #DB파일에 연결. 파일이 없으면 생성.
        # 2. SQL 테이블에 저장
        # 이 예제는 실행할 때마다 최신 데이터로 테이블을 교체.
        measurements.to_sql(    #테이블 생성과 데이터 삽입을 코드로 작성하는 대신 to_sql()이 DataFrame을 SQL 데이터베이스에 테이블로 저장한다.
            "measurements", conn, if_exists="replace", index=False  #테이블 이름은 measurements. 이미 존재하는 경우에는 덮어쓰고(replace), 인덱스는 저장되지 않도록(False)한다.
        )
        criteria.to_sql(
            "criteria", conn, if_exists="replace", index=False
        )

        # 3. SQL로 저장된 데이터 조회
        data = pd.read_sql_query(
            'SELECT * FROM measurements ORDER BY "제품번호"',
            conn,
        )
        rules = pd.read_sql_query("SELECT * FROM criteria", conn)

        summaries = []
        details = []

        # 4. 제품별로 Python 조건문을 사용해 판정
        notify("제품별 검사")
        for _, product in data.iterrows(): #iterrows() : Pandas 데이터프레임의 모든 행을 하나씩 반복문으로 꺼내서 (인덱스, 시리즈데이터)형태의 튜플로 반환함.
            failures = []
            unresolved = []

            for item in items:
                # 공통 기준 또는 해당 제품 소재의 기준을 찾는다.
                matched = rules[
                    (rules["검사항목"] == item)
                    & (
                        (rules["적용소재"] == "공통")
                        | (rules["적용소재"] == product["소재종류"])
                    )
                ]

                value = product[item]
                lower = upper = None

                if len(matched) != 1:
                    status = "판정불가"
                    reason = f"{item}: 적용 기준 누락 또는 중복"
                else:
                    rule = matched.iloc[0]
                    lower = rule["하한"]
                    upper = rule["상한"]
                    status, reason = check_value(
                        item, value, lower, upper
                    )

                if status == "불합격":
                    failures.append(reason)
                elif status == "판정불가":
                    unresolved.append(reason)

                details.append({
                    "제품번호": product["제품번호"],
                    "검사항목": item,
                    "측정값": value,
                    "하한": lower,
                    "상한": upper,
                    "항목판정": status,
                    "사유": reason,
                })

            if failures:
                final_status = "불합격"
            elif unresolved:
                final_status = "판정불가"
            else:
                final_status = "합격"

            reasons = failures + unresolved

            summaries.append({
                **product.to_dict(),
                "판정": final_status,
                "판정사유": (
                    "\n".join(reasons)
                    if reasons
                    else f"{len(items)}개 항목 모두 허용범위 이내"
                ),
            })

        result = pd.DataFrame(summaries)
        detail = pd.DataFrame(details)

        # 5. 판정 결과도 SQL에 저장
        notify("검사 결과 DB 저장")
        result.to_sql(
            "inspection_results", conn,
            if_exists="replace", index=False
        )
        detail.to_sql(
            "inspection_details", conn,
            if_exists="replace", index=False
        )

    status_counts = result["판정"].value_counts()
    report = InspectionReport(
        measurements=measurements,
        criteria=criteria,
        result=result,
        detail=detail,
        counts={
            "전체": len(result),
            **{status: int(status_counts.get(status, 0)) for status in ("합격", "불합격", "판정불가")},
        },
        input_path=input_path,
        db_path=db_path,
        excel_path=excel_path,
    )

    # 6. 결과 Excel 생성
    # 저장 실패 시에도 검사 데이터는 report에 남습니다.
    notify("결과 Excel 저장")
    save_result_excel(report)
    notify("완료" if report.excel_saved else "DB 저장 완료 / Excel 저장 실패")
    return report


#터미널에서 실행 시 검사시작&결과요약
def main():
    """기존 터미널 실행에서는 기본 경로를 사용하고 요약을 출력합니다."""
    report = run_inspection()
    print(report.result["판정"].value_counts().to_string())
    print(f"\nDB 저장: {report.db_path}")
    if not report.excel_saved:
        print(f"Excel 저장 실패: {report.excel_path}")
        print(f"사유: {report.excel_error}")
        print("검사와 DB 저장은 완료됐습니다. 파일 열림 여부와 저장 권한을 확인하세요.")
        raise SystemExit(1)

    print(f"Excel 저장: {report.excel_path}")


if __name__ == "__main__":
    main()


